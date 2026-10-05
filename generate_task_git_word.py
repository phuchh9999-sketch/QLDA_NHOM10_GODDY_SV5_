# -*- coding: utf-8 -*-
"""
Script tạo tài liệu Word (.docx) chuyên nghiệp:
BẢNG ÁNH XẠ DANH MỤC TASK VÀ FILE GIT COMMIT LÊN GITHUB
Dự án: GODDY Recruit - Quản lý Hóa đơn / Công nợ Tuyển dụng (Nhóm 10 - HUTECH)
"""

import sys
sys.stdout.reconfigure(encoding='utf-8')
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Đặt màu nền cho cell trong bảng"""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    """Đặt padding cho ô (dxa: 20 dxa = 1 pt)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1"):
    """Đặt viền thanh lịch cho bảng"""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="1E3A8A"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_phuoc_word_document():
    doc = docx.Document()

    # Cấu hình lề trang A4
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
        
        # Footer
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("QLDA Nhóm 10 - GODDY Recruit | Ánh Xạ Task & File Git Lên GitHub")
        f_run.font.name = "Calibri"
        f_run.font.size = Pt(8.5)
        f_run.font.italic = True
        f_run.font.color.rgb = RGBColor(148, 163, 184)

    # 1. TIÊU ĐỀ BÁO CÁO
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_after = Pt(2)
    r_inst = p_inst.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ TP.HCM (HUTECH)\nKHOA CÔNG NGHỆ THÔNG TIN - BỘ MÔN CÔNG NGHỆ PHẦN MỀM")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(10)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(71, 85, 105)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("BẢNG ÁNH XẠ DANH MỤC TASK & FILE GIT COMMIT")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 58, 138)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("Dự án: Phần Mềm Quản Lý Hóa Đơn / Công Nợ Tuyển Dụng (GODDY Recruit)\nGVHD: ThS. Nguyễn Hữu Trung | Đơn vị: Nhóm 10 - Lớp 23DTHC3")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    # THÔNG TIN THÀNH VIÊN
    p_mem = doc.add_paragraph()
    p_mem.paragraph_format.space_after = Pt(10)
    set_cell_background_box = p_mem.add_run(
        "★ Thành viên Frontend Developer: Nguyễn Hoàng Phước (MSSV: 2380601770) - 47 Tasks\n"
        "★ Thành viên Database & QA: Huỳnh Nguyễn Vĩnh Phúc (MSSV: 2380614923) - 41 Tasks\n"
        "★ Quy tắc Commit: git commit -m \"TenTask_tenthanhvien\" (Đúng chuẩn WBS 220 Tasks)"
    )
    set_cell_background_box.font.name = "Calibri"
    set_cell_background_box.font.size = Pt(9.5)
    set_cell_background_box.font.bold = True
    set_cell_background_box.font.color.rgb = RGBColor(15, 23, 42)

    # ==========================================================================
    # PHẦN 1: BẢNG 47 TASK CỦA NGUYỄN HOÀNG PHƯỚC (FE)
    # ==========================================================================
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(10)
    p_h1.paragraph_format.space_after = Pt(4)
    r_h1 = p_h1.add_run("PHẦN 1: DANH MỤC 47 TASK & FILE GIT - NGUYỄN HOÀNG PHƯỚC (FE)")
    r_h1.font.name = "Calibri"
    r_h1.font.size = Pt(13)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(30, 58, 138)

    phuoc_tasks_data = [
        # Tuần 1
        ("TASK-019", "Tuần 1", "Phác thảo Wireframe Prototype UI trên Figma", "ThietKeWireframeFigmaUI_nguyenhoangphuoc", "docs/ hoặc public/images/wireframe/"),
        ("TASK-020", "Tuần 1", "Xây dựng Design System, biến CSS Token màu Slate/Indigo", "XayDungDesignSystemCssToken_nguyenhoangphuoc", "public/css/style.css"),
        ("TASK-021", "Tuần 1", "Dựng khung bố cục dùng chung Sidebar + Top Header", "DungKhungBoCucSidebarHeader_nguyenhoangphuoc", "public/css/style.css, public/js/common.js"),
        ("TASK-022", "Tuần 1", "Tạo trang khung Dashboard dữ liệu skeleton index.html", "TaoTrangDashboardSkeleton_nguyenhoangphuoc", "public/index.html"),
        ("TASK-023", "Tuần 1", "Tạo trang khung Quản lý Khách hàng skeleton clients.html", "TaoTrangKhachHangSkeleton_nguyenhoangphuoc", "public/clients.html"),
        ("TASK-024", "Tuần 1", "Tích hợp icon FontAwesome / Lucide và Font Google Inter", "TichHopFontVaIcons_nguyenhoangphuoc", "public/css/style.css, public/index.html"),
        
        # Tuần 2
        ("TASK-050", "Tuần 2", "Thiết kế màn hình Đăng nhập login.html lưu JWT localStorage", "ThietKeGiaoDienDangNhap_nguyenhoangphuoc", "public/login.html"),
        ("TASK-051", "Tuần 2", "Cơ chế phân quyền giao diện: Tự động ẩn/hiện menu RBAC", "PhanQuyenGiaoDienTheoRole_nguyenhoangphuoc", "public/js/common.js"),
        ("TASK-052", "Tuần 2", "Xây dựng Bảng danh sách Khách hàng B2B & tìm kiếm realtime", "ThietKeBangDanhSachKhachHangUI_nguyenhoangphuoc", "public/clients.html, public/js/clients.js"),
        ("TASK-053", "Tuần 2", "Thiết kế Modal Thêm mới Đối tác Doanh nghiệp B2B", "ThietKeModalThemKhachHangMoiUI_nguyenhoangphuoc", "public/clients.html (#addClientModal)"),
        ("TASK-054", "Tuần 2", "Xây dựng chức năng Xuất danh sách khách hàng ra file CSV", "XuatDanhSachKhachHangRaCSV_nguyenhoangphuoc", "public/js/clients.js (exportClientsToCSV)"),
        ("TASK-055", "Tuần 2", "Xây dựng Bảng danh sách Deal và Huy hiệu bảo hành", "ThietKeBangDanhSachDealUI_nguyenhoangphuoc", "public/recruitment.html"),
        ("TASK-056", "Tuần 2", "Thiết kế Modal Chốt Deal: Tải Job & ước tính hoa hồng", "ThietKeModalChotDealPlacementUI_nguyenhoangphuoc", "public/recruitment.html (#addPlacementModal)"),
        
        # Tuần 3
        ("TASK-082", "Tuần 3", "Xây dựng Giao diện Bảng quản lý hóa đơn invoices.html", "ThietKeBangDanhSachHoaDonUI_nguyenhoangphuoc", "public/invoices.html"),
        ("TASK-083", "Tuần 3", "Thiết kế Huy hiệu trạng thái HĐ (Sent, Partial, Paid, Overdue)", "HienThiHuyHieuTrangThaiHoaDon_nguyenhoangphuoc", "public/invoices.html, public/css/style.css"),
        ("TASK-084", "Tuần 3", "Thiết kế Modal xem chi tiết hóa đơn theo mẫu GTGT chuẩn", "ThietKeModalXemChiTietHoaDonUI_nguyenhoangphuoc", "public/invoices.html (#viewInvoiceModal)"),
        ("TASK-085", "Tuần 3", "Viết CSS @media print tối ưu hóa hiển thị khi in ấn A4", "MauInHoaDonDienTuStandard_nguyenhoangphuoc", "public/css/style.css (@media print)"),
        ("TASK-086", "Tuần 3", "Tích hợp tính năng In trực tiếp và Xuất PDF window.print()", "ChucNangInHoaDonVaXuatPDF_nguyenhoangphuoc", "public/invoices.html (printInvoice)"),
        ("TASK-087", "Tuần 3", "Thêm Bộ lọc hóa đơn theo trạng thái (Sent/Partial/Paid...)", "BoLocHoaDonTheoTrangThaiUI_nguyenhoangphuoc", "public/invoices.html (filterInvoices)"),
        ("TASK-088", "Tuần 3", "Xây dựng chức năng Xuất toàn bộ hóa đơn ra file CSV", "XuatDanhSachHoaDonRaCSV_nguyenhoangphuoc", "public/invoices.html (exportInvoicesToCSV)"),

        # Tuần 4
        ("TASK-114", "Tuần 4", "Thiết kế Giao diện Quản lý Công nợ debt.html đồng bộ", "ThietKeGiaoDienQuanLyCongNoUI_nguyenhoangphuoc", "public/debt.html"),
        ("TASK-115", "Tuần 4", "Thiết kế 3 Thẻ chỉ số Tuổi nợ Aging: 1-30, 31-60, >60 ngày", "ThietKe3CardTuoiNoAgingUI_nguyenhoangphuoc", "public/debt.html (3 cards aging)"),
        ("TASK-116", "Tuần 4", "Xây dựng Bảng chi tiết nợ quá hạn và phân loại màu", "ThietKeBangChiTietCongNoUI_nguyenhoangphuoc", "public/debt.html (table overdue)"),
        ("TASK-117", "Tuần 4", "Thiết kế Modal Thu tiền nợ: Hiển thị nợ & nhập tiền trả", "ThietKeModalThuTienTraNoUI_nguyenhoangphuoc", "public/debt.html (#payDebtModal)"),
        ("TASK-118", "Tuần 4", "Tích hợp Nút Nhắc nợ trực tiếp trên từng dòng & Toast", "NutNhacNoTuDongTrenTungDong_nguyenhoangphuoc", "public/debt.html (sendReminderToast)"),
        ("TASK-119", "Tuần 4", "Xây dựng chức năng Xuất báo cáo Tuổi nợ Aging ra file CSV", "XuatBaoCaoTuoiNoRaCSV_nguyenhoangphuoc", "public/debt.html (exportAgingToCSV)"),
        ("TASK-120", "Tuần 4", "Chuẩn bị khung hiển thị 4 Card KPI và 2 biểu đồ Dashboard", "ChuanBiKhungDashboardKPI_nguyenhoangphuoc", "public/index.html (chart containers)"),

        # Tuần 5
        ("TASK-146", "Tuần 5", "Thiết kế 4 Thẻ KPI Dashboard với hiệu ứng gradient", "ThietKe4CardKPIDashboardUI_nguyenhoangphuoc", "public/index.html (4 stats cards)"),
        ("TASK-147", "Tuần 5", "Thiết kế 2 Thanh tiến trình Tỷ lệ thu hồi & Tỷ lệ nợ xấu", "ThietKe2TheTyLeThuHoiVaNoXau_nguyenhoangphuoc", "public/index.html (progress bars)"),
        ("TASK-148", "Tuần 5", "Tích hợp Chart.js vẽ Biểu đồ đường Doanh thu 12 tháng", "VeBieuDoDuongDoanhThuTheoThang_nguyenhoangphuoc", "public/js/dashboard.js (renderLineChart)"),
        ("TASK-149", "Tuần 5", "Vẽ Biểu đồ tròn Doughnut Chart Cơ cấu ngành nghề đối tác", "VeBieuDoTronCoCauNganhNghe_nguyenhoangphuoc", "public/js/dashboard.js (renderDoughnutChart)"),
        ("TASK-150", "Tuần 5", "Thiết kế Bảng Top 5 Con Nợ lớn nhất kèm nút xem chi tiết", "ThietKeBangTop5DebtorsTrenDashboard_nguyenhoangphuoc", "public/index.html, public/js/dashboard.js"),
        ("TASK-151", "Tuần 5", "Tạo Nút Làm mới Dữ liệu Dashboard với hiệu ứng loading", "NutLamMoiDuLieuDashboard_nguyenhoangphuoc", "public/index.html (refreshDashboard)"),
        ("TASK-152", "Tuần 5", "Tối ưu hóa hiển thị giao diện Dashboard Responsive di động", "ToiUuGiaoDienDashboardResponsive_nguyenhoangphuoc", "public/css/style.css (responsive styles)"),

        # Tuần 6
        ("TASK-177", "Tuần 6", "Xây dựng Giao diện Trang Nhật ký Audit Log audit.html", "ThietKeGiaoDienBangAuditLogUI_nguyenhoangphuoc", "public/audit.html"),
        ("TASK-178", "Tuần 6", "Thiết kế huy hiệu màu sắc cho các hành động Audit Log", "HienThiBadgeHanhDongAuditLog_nguyenhoangphuoc", "public/audit.html, public/css/style.css"),
        ("TASK-179", "Tuần 6", "Xây dựng Giao diện Cổng tra cứu Khách hàng portal.html", "ThietKeGiaoDienClientPortal_nguyenhoangphuoc", "public/client-portal.html, public/js/client-portal.js"),
        ("TASK-180", "Tuần 6", "Tối ưu hóa Micro-interactions và thông báo Toast trên UI", "ToiUuHieuUngToastVaMicroInteractions_nguyenhoangphuoc", "public/css/style.css, public/js/common.js"),
        ("TASK-181", "Tuần 6", "Kiểm thử khả năng tương thích Chrome, Edge, Firefox", "KiemTraTuongThichTrinhDuyetWeb_nguyenhoangphuoc", "docs/ (Báo cáo Cross-Browser Test)"),
        ("TASK-182", "Tuần 6", "Tối ưu dung lượng CSS, JS giảm thời gian tải trang < 1s", "ToiUuDungLuongFrontendAsset_nguyenhoangphuoc", "public/css/style.css, public/js/"),
        ("TASK-183", "Tuần 6", "Chụp ảnh màn hình các phân hệ giao diện cho User Manual", "ChupAnhManHinhGiaoDienUserManual_nguyenhoangphuoc", "public/images/screens/"),
        ("TASK-215", "Tuần 6", "Tạo Modal Cập nhật Doanh nghiệp B2B & Credit Limit", "ThietKeModalChinhSuaKhachHangUI_nguyenhoangphuoc", "public/clients.html (#editClientModal)"),
        ("TASK-216", "Tuần 6", "Thêm Bộ lọc thời gian (Tháng, Quý, Năm) cho Dashboard", "BoLocDashboardTheoThangQuyNam_nguyenhoangphuoc", "public/index.html, public/js/dashboard.js"),

        # Tuần 7
        ("TASK-204", "Tuần 7", "Kiểm tra toàn bộ liên kết điều hướng và tính đồng nhất 5 trang", "KiemTraLienKetDieuHuongDongNhat_nguyenhoangphuoc", "public/*.html (Kiểm tra href / menu)"),
        ("TASK-205", "Tuần 7", "Tinh chỉnh thông báo lỗi và trải nghiệm người dùng cuối", "TinhChinhThongBaoLoiNguoiDung_nguyenhoangphuoc", "public/js/common.js (Toast messages)"),
        ("TASK-206", "Tuần 7", "Hỗ trợ chuẩn bị hình ảnh minh họa cho Slide bảo vệ", "HoTroHinhAnhSlideBaoVe_nguyenhoangphuoc", "docs/slides/ (assets minh họa UI)"),
        ("TASK-207", "Tuần 7", "Thực hiện demo trực tiếp các luồng giao diện trước hội đồng", "DemoGiaoDienTrucTiepHoiDong_nguyenhoangphuoc", "Kịch bản demo UI thực tế")
    ]

    # Dựng bảng
    table_p = doc.add_table(rows=1, cols=5)
    table_p.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_p.autofit = False
    set_table_borders(table_p)

    headers = ["Mã Task", "Tuần", "Tên Công Việc Chi Tiết", "Quy Ước Commit Git", "File Cần git add"]
    col_widths = [Inches(0.9), Inches(0.7), Inches(2.3), Inches(1.8), Inches(1.8)]

    hdr_row = table_p.rows[0]
    for idx, name in enumerate(headers):
        cell = hdr_row.cells[idx]
        cell.text = name
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.runs[0]
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for idx, (t_id, t_week, t_name, t_commit, t_file) in enumerate(phuoc_tasks_data):
        row = table_p.add_row()
        bg_color = "F8FAFC" if idx % 2 == 1 else "FFFFFF"

        vals = [t_id, t_week, t_name, t_commit, t_file]
        for c_idx, val in enumerate(vals):
            cell = row.cells[c_idx]
            cell.text = val
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=60, bottom=60, left=70, right=70)
            p = cell.paragraphs[0]
            if c_idx in [0, 1]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.runs[0]
            r.font.name = "Calibri"
            r.font.size = Pt(8)
            if c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(30, 58, 138)
            elif c_idx == 3:
                r.font.name = "Consolas"
                r.font.size = Pt(7.5)
                r.font.color.rgb = RGBColor(15, 23, 42)
            elif c_idx == 4:
                r.font.bold = True
                r.font.color.rgb = RGBColor(3, 105, 161)

    for row in table_p.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    # ==========================================================================
    # PHẦN 2: BẢNG 41 TASK CỦA HUỲNH NGUYỄN VĨNH PHÚC (DBA & QA - SV5)
    # ==========================================================================
    p_h2 = doc.add_paragraph()
    p_h2.paragraph_format.space_before = Pt(20)
    p_h2.paragraph_format.space_after = Pt(4)
    r_h2 = p_h2.add_run("PHẦN 2: DANH MỤC 41 TASK & FILE GIT - HUỲNH NGUYỄN VĨNH PHÚC (DBA & QA)")
    r_h2.font.name = "Calibri"
    r_h2.font.size = Pt(13)
    r_h2.font.bold = True
    r_h2.font.color.rgb = RGBColor(30, 58, 138)

    phuc_tasks_data = [
        # Tuần 1
        ("TASK-025", "Tuần 1", "Thiết kế Sơ đồ quan hệ thực thể ERD 8 thực thể", "ThietKeSoDoQuanHeERD_huynhnguyenvinhphuc", "ERD_QLDA_NHOM10_GODDY.drawio"),
        ("TASK-026", "Tuần 1", "Viết script DDL PostgreSQL và SQL Server chuẩn khóa ngoại", "VietScriptDDLPostgresSqlserver_huynhnguyenvinhphuc", "database_supabase_postgres.sql, database_sqlserver.sql"),
        ("TASK-027", "Tuần 1", "Xây dựng Model User (Sequelize ORM) phân quyền 3 role", "ThietKeModelUser_huynhnguyenvinhphuc", "models/User.js"),
        ("TASK-028", "Tuần 1", "Xây dựng Model Client (Khách hàng B2B, Net days, Limit)", "ThietKeModelClient_huynhnguyenvinhphuc", "models/Client.js"),
        ("TASK-029", "Tuần 1", "Viết script nạp dữ liệu mẫu Enterprise (FPT, VNG, Shopee)", "VietScriptNapDuLieuMauEnterprise_huynhnguyenvinhphuc", "scripts/seed_enterprise_data.js"),
        ("TASK-030", "Tuần 1", "Cài đặt testing framework và viết Unit Test đồng bộ CSDL", "CaiDatTestingJestKiemTraDB_huynhnguyenvinhphuc", "tests/unit.test.js"),

        # Tuần 2
        ("TASK-057", "Tuần 2", "Xây dựng Model Job và Model Candidate ràng buộc ClientId", "ThietKeModelJobVaCandidate_huynhnguyenvinhphuc", "models/Job.js, models/Candidate.js"),
        ("TASK-058", "Tuần 2", "Xây dựng Model Placement ghi nhận thỏa thuận tuyển dụng", "ThietKeModelPlacement_huynhnguyenvinhphuc", "models/Placement.js"),
        ("TASK-059", "Tuần 2", "Kiểm thử tự động Xác thực Đăng nhập & cấp Token JWT", "BoKiemThuTuDong_AuthJWT_huynhnguyenvinhphuc", "tests/auth_jwt.test.js"),
        ("TASK-060", "Tuần 2", "Kiểm thử tự động Thêm khách hàng, chặn trùng MST", "BoKiemThuTuDong_ThemKhachHang_huynhnguyenvinhphuc", "tests/client_validation.test.js"),
        ("TASK-061", "Tuần 2", "Kiểm thử tự động Chốt Deal & tính ngày hết hạn bảo hành", "BoKiemThuTuDong_ChotDealPlacement_huynhnguyenvinhphuc", "tests/placement_deal.test.js"),
        ("TASK-062", "Tuần 2", "Đo lường độ bao phủ Code Coverage Sprint 1 đạt > 80%", "DoLuongCodeCoverageSprint1_huynhnguyenvinhphuc", "docs/BAO_CAO_COVERAGE_SPRINT1.md, tests/coverage_sprint1_report.js"),

        # Tuần 3
        ("TASK-089", "Tuần 3", "Xây dựng Model Invoice (subtotal, vatRate, vatAmount...)", "ThietKeModelInvoice_huynhnguyenvinhphuc", "models/Invoice.js"),
        ("TASK-090", "Tuần 3", "Cấu hình quan hệ Placement 1 - 1 Invoice trong Sequelize", "ThietLapQuanHePlacementVaInvoice_huynhnguyenvinhphuc", "tests/placement_invoice_relation.test.js"),
        ("TASK-091", "Tuần 3", "Kiểm thử tự động Tính thuế VAT 8% và Tổng tiền hóa đơn", "BoKiemThuTuDong_PhatHanhHoaDon_huynhnguyenvinhphuc", "tests/vat_calculation.test.js"),
        ("TASK-092", "Tuần 3", "Kiểm thử tự động Chặn phát hành trùng Hóa đơn trên 1 Deal", "BoKiemThuTuDong_ChanTrungHoaDon_huynhnguyenvinhphuc", "tests/duplicate_invoice_prevention.test.js"),
        ("TASK-093", "Tuần 3", "Cập nhật seed data mẫu với các hóa đơn thực tế đa trạng thái", "CapNhatSeedDataHoaDonMau_huynhnguyenvinhphuc", "scripts/seed_invoices_data.js"),
        ("TASK-094", "Tuần 3", "Đối soát toàn vẹn Placement.serviceFee và Invoice.subtotal", "KiemTraToanVenDuLieuFeeVaSubtotal_huynhnguyenvinhphuc", "tests/reconciliation_fee_subtotal.test.js"),

        # Tuần 4
        ("TASK-121", "Tuần 4", "Xây dựng Model Payment (amount, paymentDate, referenceNo)", "ThietKeModelPayment_huynhnguyenvinhphuc", "models/Payment.js"),
        ("TASK-122", "Tuần 4", "Cấu hình quan hệ Invoice 1 - N Payment trong Sequelize", "ThietLapQuanHeInvoiceVaPayment_huynhnguyenvinhphuc", "tests/invoice_payment_relation.test.js"),
        ("TASK-123", "Tuần 4", "Kiểm thử tự động Thanh toán trừ dần nợ & chống trả thừa", "BoKiemThuTuDong_ThanhToanTruNo_huynhnguyenvinhphuc", "tests/debt_reduction_payment.test.js"),
        ("TASK-124", "Tuần 4", "Kiểm thử Thuật toán phân loại Tuổi nợ Aging 3 xô", "BoKiemThuTuDong_PhanTichTuoiNo_huynhnguyenvinhphuc", "tests/aging_analysis.test.js"),
        ("TASK-125", "Tuần 4", "Tạo chỉ mục Database Indexes tối ưu tốc độ dueDate, status", "TaoIndexToiUuTruyVanDueDate_huynhnguyenvinhphuc", "scripts/migrate_optimize_indexes.js"),
        ("TASK-126", "Tuần 4", "Lập báo cáo nghiệm thu Acceptance Criteria Sprint 2", "BaoCaoKiemThuSprint2_huynhnguyenvinhphuc", "docs/BAO_CAO_KIEM_THU_QA_SPRINT2.md, tests/sprint2_acceptance.test.js"),

        # Tuần 5
        ("TASK-153", "Tuần 5", "Kiểm thử tự động tính toán số liệu thống kê Dashboard KPI", "BoKiemThuTuDong_DashboardStats_huynhnguyenvinhphuc", "tests/dashboard_stats.test.js"),
        ("TASK-154", "Tuần 5", "Đối soát mảng doanh thu 12 tháng khớp 100% bản ghi Payment", "KiemThuDoiSoatDoanhThu12Thang_huynhnguyenvinhphuc", "tests/reconciliation_revenue_12months.test.js"),
        ("TASK-155", "Tuần 5", "Đối chiếu danh sách Top 5 Debtors khớp hóa đơn quá hạn", "KiemThuTop5DebtorsKhopHoaDon_huynhnguyenvinhphuc", "tests/top5_debtors.test.js"),
        ("TASK-156", "Tuần 5", "Đo lường thời gian phản hồi API Dashboard SLA < 500ms (6.26ms)", "DoLuongThoiGianPhanHoiDashboard_huynhnguyenvinhphuc", "tests/dashboard_performance.test.js"),
        ("TASK-157", "Tuần 5", "Tổng hợp Báo cáo kiểm thử chất lượng Sprint 3", "TongHopBaoCaoKiemThuSprint3_huynhnguyenvinhphuc", "docs/BAO_CAO_KIEM_THU_QA_SPRINT3.md"),

        # Tuần 6
        ("TASK-184", "Tuần 6", "Kiểm thử tự động Ghi nhật ký Audit Log đầy đủ các trường", "BoKiemThuTuDong_AuditLog_huynhnguyenvinhphuc", "tests/audit_log_verification.test.js"),
        ("TASK-185", "Tuần 6", "Xây dựng kịch bản kiểm thử tự động toàn trình End-to-End", "KiemThuTuDongToanTrinhE2E_huynhnguyenvinhphuc", "tests/e2e_fullflow.test.js"),
        ("TASK-186", "Tuần 6", "Kiểm thử bảo mật API: Bcrypt hashing & phân quyền RBAC", "KiemThuBaoMatAPI_RBAC_huynhnguyenvinhphuc", "tests/security_rbac.test.js"),
        ("TASK-187", "Tuần 6", "Phân tích biểu đồ Pareto nguyên nhân trễ hạn công nợ 80/20", "PhanTichParetoRuiRoVaNguyenNhanNo_huynhnguyenvinhphuc", "HuynhNguyenVinhPhuc_Pareto.xlsx"),
        ("TASK-188", "Tuần 6", "Thiết lập script sao lưu và phục hồi CSDL tự động Snapshot", "ThietLapScriptBackupRestoreDB_huynhnguyenvinhphuc", "scripts/backup_restore.js"),
        ("TASK-189", "Tuần 6", "Cấu hình lệnh npm test chạy 20 test suites, Coverage > 85%", "CauHinhLenhNpmTestToanBo_huynhnguyenvinhphuc", "package.json, tests/run_all_tests.js"),
        ("TASK-218", "Tuần 6", "Kiểm thử cảnh báo 3 cấp khi khách hàng vượt Credit Limit", "KiemThuCanhBaoVuotHanMucNo_huynhnguyenvinhphuc", "tests/credit_limit_guard.test.js"),

        # Tuần 7
        ("TASK-208", "Tuần 7", "Chạy lần cuối toàn bộ test suite npm test pass 100%", "ChayFinalTestSuiteNpmTest_huynhnguyenvinhphuc", "tests/run_all_tests.js"),
        ("TASK-209", "Tuần 7", "Kiểm tra tính toàn vẹn CSDL sau kịch bản demo (8/8 PASS)", "KiemTraToanVenCSDLSauDemo_huynhnguyenvinhphuc", "scripts/check_db_integrity.js"),
        ("TASK-210", "Tuần 7", "Hoàn thiện báo cáo phân tích Pareto kiểm soát chất lượng", "HoanThienBangPhanTichParetoDocx_huynhnguyenvinhphuc", "BAO_CAO_PHAN_TICH_PARETO_HUYNHNGUYENVINHPHUC.docx"),
        ("TASK-211", "Tuần 7", "Đóng gói toàn bộ tài liệu Test Case, Log & Summary Report", "DongGoiBoTaiLieuKiemThuQA_huynhnguyenvinhphuc", "HO_SO_KIEM_THU_QA_TEST_REPORT.md"),
        ("TASK-212", "Tuần 7", "Đề cương ôn tập & giải trình CSDL, quy trình QA bảo vệ", "TraLoiCauHoiCSDLVaQA_huynhnguyenvinhphuc", "DE_CUONG_ON_TAP_BAO_VE_CSDL_QA.md")
    ]

    table_phuc = doc.add_table(rows=1, cols=5)
    table_phuc.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_phuc.autofit = False
    set_table_borders(table_phuc)

    hdr_row2 = table_phuc.rows[0]
    for idx, name in enumerate(headers):
        cell = hdr_row2.cells[idx]
        cell.text = name
        set_cell_background(cell, "047857") # Màu xanh ngọc Emerald
        set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.runs[0]
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for idx, (t_id, t_week, t_name, t_commit, t_file) in enumerate(phuc_tasks_data):
        row = table_phuc.add_row()
        bg_color = "F0FDF4" if idx % 2 == 1 else "FFFFFF"

        vals = [t_id, t_week, t_name, t_commit, t_file]
        for c_idx, val in enumerate(vals):
            cell = row.cells[c_idx]
            cell.text = val
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=60, bottom=60, left=70, right=70)
            p = cell.paragraphs[0]
            if c_idx in [0, 1]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.runs[0]
            r.font.name = "Calibri"
            r.font.size = Pt(8)
            if c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(4, 120, 87)
            elif c_idx == 3:
                r.font.name = "Consolas"
                r.font.size = Pt(7.5)
                r.font.color.rgb = RGBColor(15, 23, 42)
            elif c_idx == 4:
                r.font.bold = True
                r.font.color.rgb = RGBColor(5, 150, 105)

    for row in table_phuc.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    # ==========================================================================
    # PHẦN 3: HƯỚNG DẪN THỰC THI CÁC LỆNH GIT
    # ==========================================================================
    p_h3 = doc.add_paragraph()
    p_h3.paragraph_format.space_before = Pt(20)
    p_h3.paragraph_format.space_after = Pt(4)
    r_h3 = p_h3.add_run("PHẦN 3: QUY TRÌNH & CÂU LỆNH GIT THỰC HÀNH CHUẨN")
    r_h3.font.name = "Calibri"
    r_h3.font.size = Pt(13)
    r_h3.font.bold = True
    r_h3.font.color.rgb = RGBColor(30, 58, 138)

    git_steps = [
        "1. Kiểm tra trạng thái mã nguồn cục bộ: git status",
        "2. Thêm file tương ứng của task vào vùng theo dõi (Staging Area): git add <đường_dẫn_file>",
        "3. Commit đúng theo cú pháp quy định: git commit -m \"TenTask_tenthanhvien\"",
        "4. Đẩy commit lên nhánh cá nhân trên GitHub: git push origin <tên_nhánh_của_bạn>",
        "5. Tạo Pull Request (PR) để Trưởng nhóm (PM Nguyễn Hữu Phúc) review và gộp (Merge) vào nhánh main/develop."
    ]

    for step in git_steps:
        p_step = doc.add_paragraph(style='List Bullet')
        p_step.paragraph_format.space_after = Pt(2)
        r_s = p_step.add_run(step)
        r_s.font.size = Pt(9.5)
        r_s.font.color.rgb = RGBColor(30, 41, 59)

    # Lưu tài liệu Word
    out_path = r'c:\Users\phuco\QLDA_NHOM10_GODDY\BANG_ANH_XA_TASK_VA_FILE_GIT_LEN_GITHUB.docx'
    doc.save(out_path)
    print(f"[THÀNH CÔNG] Đã tạo file Word hoàn chỉnh tại: {out_path}")

    # Đồng thời tạo bản dành riêng cho Nguyễn Hoàng Phước
    doc_phuoc = docx.Document()
    for section in doc_phuoc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("Frontend Developer: Nguyễn Hoàng Phước | 47 Tasks & File Git")
        f_run.font.name = "Calibri"
        f_run.font.size = Pt(8.5)
        f_run.font.italic = True
        f_run.font.color.rgb = RGBColor(148, 163, 184)

    p_t = doc_phuoc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_t.add_run("DANH MỤC 47 TASK & FILE GIT COMMIT\nNGUYỄN HOÀNG PHƯỚC - FRONTEND DEVELOPER")
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(16)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(30, 58, 138)

    p_info = doc_phuoc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_info.paragraph_format.space_after = Pt(12)
    r_info = p_info.add_run("MSSV: 2380601770 | Lớp: 23DTHC3 | Khoa CNTT - Đại học HUTECH\nĐề tài: GODDY Recruit - Quản Lý Hóa Đơn & Công Nợ Tuyển Dụng (Nhóm 10)")
    r_info.font.size = Pt(10)
    r_info.font.color.rgb = RGBColor(71, 85, 105)

    # Thêm bảng cho Phước
    t_single = doc_phuoc.add_table(rows=1, cols=5)
    t_single.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_single.autofit = False
    set_table_borders(t_single)

    hdr = t_single.rows[0]
    for idx, name in enumerate(headers):
        cell = hdr.cells[idx]
        cell.text = name
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.runs[0]
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for idx, (t_id, t_week, t_name, t_commit, t_file) in enumerate(phuoc_tasks_data):
        row = t_single.add_row()
        bg_color = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        vals = [t_id, t_week, t_name, t_commit, t_file]
        for c_idx, val in enumerate(vals):
            cell = row.cells[c_idx]
            cell.text = val
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=60, bottom=60, left=70, right=70)
            p = cell.paragraphs[0]
            if c_idx in [0, 1]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.runs[0]
            r.font.name = "Calibri"
            r.font.size = Pt(8)
            if c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(30, 58, 138)
            elif c_idx == 3:
                r.font.name = "Consolas"
                r.font.size = Pt(7.5)
                r.font.color.rgb = RGBColor(15, 23, 42)
            elif c_idx == 4:
                r.font.bold = True
                r.font.color.rgb = RGBColor(3, 105, 161)

    for row in t_single.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    out_phuoc = r'c:\Users\phuco\QLDA_NHOM10_GODDY\DANH_SACH_TASK_VA_FILE_GIT_NGUYEN_HOANG_PHUOC.docx'
    doc_phuoc.save(out_phuoc)
    print(f"[THÀNH CÔNG] Đã tạo file Word riêng cho bạn Phước tại: {out_phuoc}")

if __name__ == '__main__':
    build_phuoc_word_document()
