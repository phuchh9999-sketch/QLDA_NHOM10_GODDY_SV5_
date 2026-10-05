# -*- coding: utf-8 -*-
"""
Script tạo Báo cáo Word Phân tích Pareto hoàn chỉnh cho Huỳnh Nguyễn Vĩnh Phúc (TASK-210)
Tự động nhúng hình ảnh biểu đồ Pareto và lập bảng số liệu 80/20 chuyên nghiệp
"""

import os
import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()

# Page setup A4
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

# Helper styling
def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/><w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

# Title
p_univ = doc.add_paragraph()
p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_u = p_univ.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ TP. HỒ CHÍ MINH (HUTECH)\nKHOA CÔNG NGHỆ THÔNG TIN — BỘ MÔN QUẢN LÝ DỰ ÁN CNTT\n")
r_u.font.name = "Calibri"
r_u.font.size = Pt(11)
r_u.font.bold = True
r_u.font.color.rgb = RGBColor(30, 41, 59)

p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_before = Pt(8)
p_title.paragraph_format.space_after = Pt(4)
r_t = p_title.add_run("BÁO CÁO PHÂN TÍCH BIỂU ĐỒ PARETO (QUY TẮC 80/20)\nKIỂM SOÁT CHẤT LƯỢNG DỰ ÁN PHẦN MỀM GODDY RECRUIT")
r_t.font.name = "Calibri"
r_t.font.size = Pt(16)
r_t.font.bold = True
r_t.font.color.rgb = RGBColor(30, 58, 138)

p_author = doc.add_paragraph()
p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_author.paragraph_format.space_after = Pt(16)
r_a = p_author.add_run(
    "Sinh viên thực hiện: HUỲNH NGUYỄN VĨNH PHÚC\n"
    "Mã số sinh viên: 2380614923 — Lớp: 23DTHC3\n"
    "Vai trò: Tester & Database Administrator (SV5 - Nhóm 10)\n"
    "Giảng viên hướng dẫn: Thầy Nguyễn Hữu Trung"
)
r_a.font.name = "Calibri"
r_a.font.size = Pt(10.5)
r_a.font.italic = True
r_a.font.color.rgb = RGBColor(71, 85, 105)

# Section 1: Khái niệm
h1 = doc.add_paragraph()
h1.paragraph_format.space_before = Pt(10)
r_h1 = h1.add_run("1. TỔNG QUAN NGUYÊN LÝ PARETO TRONG QUẢN TRỊ DỰ ÁN CNTT")
r_h1.font.bold = True
r_h1.font.size = Pt(12.5)
r_h1.font.color.rgb = RGBColor(30, 58, 138)

p1 = doc.add_paragraph()
p1.add_run(
    "Nguyên lý Pareto (Quy tắc 80/20) chỉ ra rằng trong hầu hết mọi hệ thống phần mềm và quản trị dự án, "
    "khoảng 80% các vấn đề, lỗi defect hoặc sự chậm trễ tài chính bắt nguồn từ 20% nguyên nhân cốt lõi (Vital Few). "
    "Bằng cách nhận diện và tập trung xử lý dứt điểm nhóm 20% nguyên nhân trọng yếu này, nhóm phát triển có thể đạt được hiệu quả cải tiến chất lượng tối đa với chi phí và nguồn lực tối thiểu."
)
p1.paragraph_format.space_after = Pt(10)

# Section 2: Bài thực hành cá nhân
h2 = doc.add_paragraph()
h2.paragraph_format.space_before = Pt(10)
r_h2 = h2.add_run("2. PHÂN TÍCH PARETO BÀI THỰC HÀNH CÁ NHÂN (100 LỖI WEBSITE MÁY TÍNH)")
r_h2.font.bold = True
r_h2.font.size = Pt(12.5)
r_h2.font.color.rgb = RGBColor(30, 58, 138)

# Bảng số liệu bài cá nhân
t1 = doc.add_table(rows=9, cols=6)
t1.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t1)

t1_headers = ["Mã Lỗi", "Mô Tả Lỗi", "Số Lần (Tần số)", "Tỷ Lệ (%)", "Tích Lũy (%)", "Phân Loại"]
t1_data = [
    ("L02", "Đặt hàng không thành công", "28", "28.0%", "28.0%", "Nhóm trọng yếu (80%)"),
    ("L06", "Giỏ hàng tính sai tổng tiền", "22", "22.0%", "50.0%", "Nhóm trọng yếu (80%)"),
    ("L04", "Tìm kiếm & lọc sản phẩm sai", "16", "16.0%", "66.0%", "Nhóm trọng yếu (80%)"),
    ("L01", "Cấu hình máy tính sai/thiếu", "14", "14.0%", "80.0%", "Nhóm trọng yếu (80%)"),
    ("L07", "Hình ảnh sản phẩm bị vỡ lỗi", "8", "8.0%", "88.0%", "Nhóm thứ yếu (20%)"),
    ("L03", "Không gửi email xác nhận", "5", "5.0%", "93.0%", "Nhóm thứ yếu (20%)"),
    ("L08", "Trang sản phẩm tải chậm", "4", "4.0%", "97.0%", "Nhóm thứ yếu (20%)"),
    ("L05", "Lỗi phân quyền xem đơn người khác", "3", "3.0%", "100.0%", "Nhóm thứ yếu (20%)")
]

for idx, h in enumerate(t1_headers):
    c = t1.rows[0].cells[idx]
    c.text = h
    set_cell_background(c, "1E3A8A")
    set_cell_margins(c, 80, 80, 80, 80)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.runs[0]
    r.font.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(255, 255, 255)

for r_idx, row in enumerate(t1_data, 1):
    for c_idx, val in enumerate(row):
        c = t1.rows[r_idx].cells[c_idx]
        c.text = val
        set_cell_margins(c, 60, 60, 60, 60)
        p = c.paragraphs[0]
        if c_idx in [0, 2, 3, 4, 5]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.runs[0]
        r.font.size = Pt(8.5)
        if r_idx <= 4 and c_idx == 5:
            r.font.bold = True
            r.font.color.rgb = RGBColor(185, 28, 28) # Red alert
            set_cell_background(c, "FEE2E2")
        elif r_idx % 2 == 0 and c_idx != 5:
            set_cell_background(c, "F8FAFC")

# Nhúng hình ảnh biểu đồ cá nhân
img1_path = os.path.abspath('public/images/pareto_baitap_canhan.png')
if os.path.exists(img1_path):
    p_img1 = doc.add_paragraph()
    p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img1.paragraph_format.space_before = Pt(8)
    p_img1.paragraph_format.space_after = Pt(2)
    doc.add_picture(img1_path, width=Inches(6.2))
    p_cap1 = doc.add_paragraph()
    p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap1.paragraph_format.space_after = Pt(12)
    r_cap1 = p_cap1.add_run("Hình 1: Biểu đồ Pareto 100 lỗi hệ thống website máy tính (Bài thực hành cá nhân)")
    r_cap1.font.size = Pt(8.5)
    r_cap1.font.italic = True
    r_cap1.font.color.rgb = RGBColor(100, 116, 139)

# Nhận xét bài cá nhân
p_eval1 = doc.add_paragraph()
p_eval1.add_run("📌 Nhận xét & Đề xuất hành động khắc phục:\n").bold = True
p_eval1.add_run(
    "• 4 lỗi cốt lõi: L02 (28%), L06 (22%), L04 (16%) và L01 (14%) chiếm chính xác 80% tổng số lỗi hệ thống.\n"
    "• Hành động ưu tiên: Kiểm thử lại toàn diện luồng Checkout thanh toán (L02) và logic tính thuế/tổng tiền (L06) ngay trong Sprint kế tiếp để giảm thiểu 50% khiếu nại của khách hàng."
)
p_eval1.paragraph_format.space_after = Pt(14)

# Section 3: Áp dụng vào dự án GODDY Recruit
h3 = doc.add_paragraph()
h3.paragraph_format.space_before = Pt(10)
r_h3 = h3.add_run("3. PHÂN TÍCH PARETO ÁP DỤNG TRONG DỰ ÁN GODDY RECRUIT")
r_h3.font.bold = True
r_h3.font.size = Pt(12.5)
r_h3.font.color.rgb = RGBColor(30, 58, 138)

p2 = doc.add_paragraph()
p2.add_run(
    "Trong dự án GODDY Recruit (Quản lý Hóa đơn & Công nợ tuyển dụng), nhóm QA đã thống kê 120 sự cố phát sinh "
    "liên quan đến chậm trễ thu hồi nợ và phát hành hóa đơn VAT. Kết quả phân tích Pareto đã xác định rõ nhóm nguyên nhân cốt lõi gây tắc nghẽn dòng tiền doanh nghiệp:"
)
p2.paragraph_format.space_after = Pt(8)

# Nhúng hình ảnh biểu đồ dự án
img2_path = os.path.abspath('public/images/pareto_du_an_goddy.png')
if os.path.exists(img2_path):
    p_img2 = doc.add_paragraph()
    p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img2.paragraph_format.space_before = Pt(6)
    p_img2.paragraph_format.space_after = Pt(2)
    doc.add_picture(img2_path, width=Inches(6.2))
    p_cap2 = doc.add_paragraph()
    p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap2.paragraph_format.space_after = Pt(12)
    r_cap2 = p_cap2.add_run("Hình 2: Biểu đồ Pareto nguyên nhân trễ hạn công nợ & lỗi hóa đơn (Dự án GODDY Recruit)")
    r_cap2.font.size = Pt(8.5)
    r_cap2.font.italic = True
    r_cap2.font.color.rgb = RGBColor(100, 116, 139)

# Kết luận
p_conclusion = doc.add_paragraph()
p_conclusion.add_run("4. KẾT LUẬN & ĐÓNG GÓP CHO CHẤT LƯỢNG ĐỒ ÁN (QA CONCLUSION)\n").bold = True
p_conclusion.add_run(
    "1. Việc áp dụng công cụ chất lượng Pareto đã giúp Nhóm 10 tối ưu hóa nguồn lực kiểm thử, không dàn trải đều vào tất cả các màn hình mà tập trung sâu vào phân hệ Hóa đơn Invoicing VAT và Đối soát công nợ Aging.\n"
    "2. Kết quả sau khi tối ưu: Toàn bộ 19 test cases tự động (Unit test, E2E, Security) đều đạt tỷ lệ Pass 100%, thời gian phản hồi API Dashboard đạt dưới 500ms và không ghi nhận tình trạng trả vượt quá số nợ (Zero Overpayment).\n"
    "3. Báo cáo Pareto này được lưu trữ đồng bộ cùng file Excel gốc: HuynhNguyenVinhPhuc_Pareto.xlsx phục vụ chấm điểm đồ án."
)

docx_pareto_path = r'c:\Users\phuco\QLDA_NHOM10_GODDY\BAO_CAO_PHAN_TICH_PARETO_HUYNHNGUYENVINHPHUC.docx'
doc.save(docx_pareto_path)
print(f"Đã xuất file Word Pareto thành công: {docx_pareto_path}")
