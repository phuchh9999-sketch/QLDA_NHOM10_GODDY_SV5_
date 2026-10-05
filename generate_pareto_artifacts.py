import os
import sys
import numpy as np
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker
import matplotlib.patches as patches
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Ensure UTF-8 output
sys.stdout.reconfigure(encoding='utf-8')
os.makedirs('public/images', exist_ok=True)

plt.rcParams['font.sans-serif'] = ['Segoe UI', 'Arial', 'DejaVu Sans']
plt.rcParams['axes.unicode_minus'] = False

# ==============================================================================
# 1. BIỂU ĐỒ BÀI THỰC HÀNH CÁ NHÂN (100 LỖI WEBSITE MÁY TÍNH)
# ==============================================================================
def generate_practice_chart():
    categories = [
        'L02 - Đặt hàng không thành công',
        'L06 - Giỏ hàng tính sai tổng tiền',
        'L04 - Tìm kiếm & lọc sai',
        'L01 - Cấu hình máy sai/thiếu',
        'L07 - Hình ảnh sản phẩm lỗi',
        'L03 - Không gửi email xác nhận',
        'L08 - Trang sản phẩm tải chậm',
        'L05 - Xem đơn người khác'
    ]
    counts = [28, 22, 16, 14, 8, 5, 4, 3]
    total = sum(counts)
    cum_pct = (np.cumsum(counts) / total) * 100

    fig, ax1 = plt.subplots(figsize=(12, 7.5), dpi=220)
    fig.patch.set_facecolor('#ffffff')
    ax1.set_facecolor('#ffffff')

    # Vẽ Banner xanh đầu trang chuẩn phong cách Slide PowerPoint
    fig.patches.extend([
        plt.Rectangle((0, 0.90), 1, 0.10, fill=True, color='#2b5b84', transform=fig.transFigure, figure=fig, zorder=1),
        plt.Rectangle((0.88, 0.90), 0.12, 0.10, fill=True, color='#17395c', transform=fig.transFigure, figure=fig, zorder=2)
    ])
    fig.text(0.06, 0.94, 'Pareto website | Bài thực hành cá nhân', fontsize=18, color='#ffffff', fontweight='bold', va='center')

    x = np.arange(len(categories))
    width = 0.48

    # Bars: Count of errors
    bars = ax1.bar(x, counts, width, color='#0f2b5c', edgecolor='#081c3e', label='Count', zorder=2)
    ax1.set_ylabel('Count of Errors', fontsize=12, fontweight='bold', color='#0f172a', labelpad=8)
    ax1.set_ylim(0, 35)
    ax1.set_yticks(np.arange(0, 36, 5))
    ax1.set_xticks(x)
    ax1.set_xticklabels(categories, rotation=30, ha='right', fontsize=9.5, fontweight='500', color='#0f172a')
    ax1.grid(axis='y', linestyle='-', color='#e2e8f0', linewidth=0.8, zorder=1)

    for bar in bars:
        yval = bar.get_height()
        ax1.text(bar.get_x() + bar.get_width()/2.0, yval + 0.6, f'{int(yval)}', ha='center', va='bottom', fontsize=9.5, fontweight='bold', color='#0f2b5c')

    # Trục phải: Cumulative %
    ax2 = ax1.twinx()
    ax2.plot(x, cum_pct, color='#d91414', marker='s', markersize=6.5, linewidth=2.0, label='Cumulative %', zorder=4)
    ax2.set_ylabel('', fontsize=11)
    ax2.set_ylim(0, 100)
    ax2.set_yticks(np.arange(0, 101, 10))
    ax2.yaxis.set_major_formatter(ticker.FormatStrFormatter('%.1f'))
    ax2.tick_params(axis='y', labelsize=9.5)

    # 80% Line ngang
    ax2.axhline(80, color='#38bdf8', linestyle='--', linewidth=1.8, zorder=3)
    ax2.text(4.8, 81.5, '80% Line', color='#0284c7', fontsize=11, fontweight='bold', ha='center')

    # Đường dóng thẳng đứng tại điểm cắt x=3 (L01, 80%)
    ax1.axvline(3, color='#38bdf8', linestyle='--', linewidth=1.5, zorder=3)

    # Annotations: "Vital Few" và "Trivial Many"
    ax1.text(0.8, 31, 'Vital Few', color='#0f172a', fontsize=13, fontweight='bold', ha='center')
    ax1.text(0.8, 29, '(4 nhóm = 80% lỗi)', color='#2563eb', fontsize=9.5, fontweight='600', ha='center')

    ax1.text(5.2, 17, 'Trivial Many', color='#475569', fontsize=13, fontweight='bold', ha='center')
    ax1.text(5.2, 15, '(4 nhóm = 20% lỗi)', color='#64748b', fontsize=9.5, fontweight='600', ha='center')

    # Tiêu đề biểu đồ
    fig.text(0.50, 0.84, 'Pareto Diagram', fontsize=16, fontweight='bold', color='#0f172a', ha='center')

    # Legend
    lines1, labels1 = ax1.get_legend_handles_labels()
    lines2, labels2 = ax2.get_legend_handles_labels()
    ax1.legend(lines1 + lines2, labels1 + labels2, loc='upper right', bbox_to_anchor=(0.98, 0.65), framealpha=0.9, fontsize=10.5)

    # Footer
    fig.text(0.88, 0.02, 'Quản lý dự án CNTT', fontsize=9.5, color='#475569', style='italic')

    # Canh lề phù hợp
    plt.subplots_adjust(top=0.81, bottom=0.18, left=0.08, right=0.92)

    output_path = 'public/images/pareto_baitap_canhan.png'
    plt.savefig(output_path, dpi=220)
    plt.close()
    print(f'-> Đã lưu biểu đồ thực hành: {output_path}')


# ==============================================================================
# 2. BIỂU ĐỒ PARETO DỰ ÁN GODDY RECRUIT
# ==============================================================================
def generate_project_chart():
    categories = [
        'Lệch thuế VAT 8% & làm tròn tiền',
        'Lệch trạng thái công nợ Paid/Partial',
        'Tính sai hạn bảo hành 60 ngày',
        'Trùng mã số thuế / thiếu contact',
        'Lỗi font UTF-8 in hóa đơn điện tử',
        'Báo cáo Aging tải chậm (>500 HĐ)',
        'Lỗi phân quyền xem hóa đơn đối tác',
        'Lệch định dạng cột khi xuất CSV'
    ]
    counts = [48, 36, 14, 8, 6, 4, 3, 1]
    total = sum(counts)
    cum_pct = (np.cumsum(counts) / total) * 100

    fig, ax1 = plt.subplots(figsize=(12, 7.5), dpi=220)
    fig.patch.set_facecolor('#ffffff')
    ax1.set_facecolor('#ffffff')

    # Vẽ Banner xanh navy cho GODDY Recruit
    fig.patches.extend([
        plt.Rectangle((0, 0.90), 1, 0.10, fill=True, color='#1e3a8a', transform=fig.transFigure, figure=fig, zorder=1),
        plt.Rectangle((0.88, 0.90), 0.12, 0.10, fill=True, color='#0f172a', transform=fig.transFigure, figure=fig, zorder=2)
    ])
    fig.text(0.06, 0.94, 'Pareto GODDY Recruit | Quản Lý Hóa Đơn & Công Nợ B2B', fontsize=17, color='#ffffff', fontweight='bold', va='center')

    x = np.arange(len(categories))
    width = 0.48

    bars = ax1.bar(x, counts, width, color='#0f2b5c', edgecolor='#081c3e', label='Count', zorder=2)
    ax1.set_ylabel('Count of Errors', fontsize=12, fontweight='bold', color='#0f172a', labelpad=8)
    ax1.set_ylim(0, 60)
    ax1.set_yticks(np.arange(0, 61, 10))
    ax1.set_xticks(x)
    ax1.set_xticklabels(categories, rotation=30, ha='right', fontsize=9.5, fontweight='500', color='#0f172a')
    ax1.grid(axis='y', linestyle='-', color='#e2e8f0', linewidth=0.8, zorder=1)

    for bar in bars:
        yval = bar.get_height()
        ax1.text(bar.get_x() + bar.get_width()/2.0, yval + 1.0, f'{int(yval)}', ha='center', va='bottom', fontsize=9.5, fontweight='bold', color='#0f2b5c')

    # Trục phải: Cumulative %
    ax2 = ax1.twinx()
    ax2.plot(x, cum_pct, color='#d91414', marker='s', markersize=6.5, linewidth=2.0, label='Cumulative %', zorder=4)
    ax2.set_ylabel('', fontsize=11)
    ax2.set_ylim(0, 100)
    ax2.set_yticks(np.arange(0, 101, 10))
    ax2.yaxis.set_major_formatter(ticker.FormatStrFormatter('%.1f'))
    ax2.tick_params(axis='y', labelsize=9.5)

    # 80% Line ngang
    ax2.axhline(80, color='#38bdf8', linestyle='--', linewidth=1.8, zorder=3)
    ax2.text(4.8, 81.5, '80% Line', color='#0284c7', fontsize=11, fontweight='bold', ha='center')

    # Đường dóng thẳng đứng tại index 2 (Hạn bảo hành 60d, đạt 81.7%)
    ax1.axvline(2, color='#38bdf8', linestyle='--', linewidth=1.5, zorder=3)

    # Annotations: "Vital Few" và "Trivial Many"
    ax1.text(0.7, 56.5, 'Vital Few', color='#0f172a', fontsize=13, fontweight='bold', ha='center')
    ax1.text(0.7, 53.5, '(3 nhóm = 81.7% lỗi)', color='#2563eb', fontsize=9.5, fontweight='600', ha='center')

    ax1.text(5.0, 30, 'Trivial Many', color='#475569', fontsize=13, fontweight='bold', ha='center')
    ax1.text(5.0, 26, '(5 nhóm = 18.3% lỗi)', color='#64748b', fontsize=9.5, fontweight='600', ha='center')

    # Tiêu đề biểu đồ
    fig.text(0.50, 0.84, 'Pareto Diagram - Chất Lượng Hệ Thống GODDY Recruit', fontsize=16, fontweight='bold', color='#0f172a', ha='center')

    # Legend
    lines1, labels1 = ax1.get_legend_handles_labels()
    lines2, labels2 = ax2.get_legend_handles_labels()
    ax1.legend(lines1 + lines2, labels1 + labels2, loc='upper right', bbox_to_anchor=(0.98, 0.65), framealpha=0.9, fontsize=10.5)

    # Footer
    fig.text(0.84, 0.02, 'Quản lý dự án CNTT - Nhóm 10', fontsize=9.5, color='#475569', style='italic')

    plt.subplots_adjust(top=0.81, bottom=0.18, left=0.08, right=0.92)

    output_path = 'public/images/pareto_du_an_goddy.png'
    plt.savefig(output_path, dpi=220)
    plt.close()
    print(f'-> Đã lưu biểu đồ dự án: {output_path}')


# ==============================================================================
# 3. TẠO FILE EXCEL HOÀN CHỈNH NỘP BÀI (MSSV_HoTen_Pareto.xlsx)
# ==============================================================================
def generate_excel_file():
    wb = openpyxl.Workbook()
    
    # --------------------------------------------------------------------------
    # SHEET 1: BÀI THỰC HÀNH CÁ NHÂN (Theo đúng đề bài yêu cầu)
    # --------------------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Thực hành cá nhân Pareto"
    ws1.views.sheetView[0].showGridLines = True

    font_title = Font(name='Segoe UI', size=14, bold=True, color='1E3A8A')
    font_subtitle = Font(name='Segoe UI', size=10.5, italic=True, color='475569')
    font_meta = Font(name='Segoe UI', size=11, bold=True, color='0F172A')
    font_meta_val = Font(name='Segoe UI', size=11, color='1E293B')
    font_sec_head = Font(name='Segoe UI', size=11.5, bold=True, color='FFFFFF')
    font_tbl_head = Font(name='Segoe UI', size=10.5, bold=True, color='FFFFFF')
    font_tbl_cell = Font(name='Segoe UI', size=10.5, color='0F172A')
    font_tbl_bold = Font(name='Segoe UI', size=10.5, bold=True, color='0F172A')
    font_q_head = Font(name='Segoe UI', size=10.5, bold=True, color='1E3A8A')
    font_ans = Font(name='Segoe UI', size=10, color='1E293B')

    fill_navy = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    fill_blue_head = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    fill_total = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')
    fill_vital = PatternFill(start_color='DBEAFE', end_color='DBEAFE', fill_type='solid')
    fill_trivial = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    fill_q_box = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')

    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    border_total = Border(
        left=Side(style='thin', color='94A3B8'),
        right=Side(style='thin', color='94A3B8'),
        top=Side(style='medium', color='1E3A8A'),
        bottom=Side(style='double', color='1E3A8A')
    )

    ws1['A1'] = "BÀI THỰC HÀNH CÁ NHÂN: PHÂN TÍCH VÀ VẼ BIỂU ĐỒ PARETO"
    ws1['A1'].font = font_title
    ws1['A2'] = "Môn: Quản lý dự án CNTT  |  Thời gian: 15 phút  |  Công cụ: Microsoft Excel"
    ws1['A2'].font = font_subtitle

    ws1['A4'] = "Họ và tên:"
    ws1['A4'].font = font_meta
    ws1['B4'] = "Huỳnh Nguyễn Vĩnh Phúc"
    ws1['B4'].font = font_meta_val

    ws1['D4'] = "MSSV:"
    ws1['D4'].font = font_meta
    ws1['E4'] = "(Điền MSSV của bạn)"
    ws1['E4'].font = font_meta_val

    ws1['F4'] = "Lớp:"
    ws1['F4'].font = font_meta
    ws1['G4'] = "Quản lý dự án CNTT - Nhóm 10"
    ws1['G4'].font = font_meta_val

    ws1['A6'] = "1. BẢNG DỮ LIỆU BÁO CÁO LỖI VÀ TÍNH TOÁN LŨY KẾ (Đã sắp xếp giảm dần theo Số báo cáo)"
    ws1.merge_cells('A6:G6')
    ws1['A6'].font = font_sec_head
    ws1['A6'].fill = fill_navy
    ws1['A6'].alignment = Alignment(vertical='center', indent=1)

    headers = [
        ('A7', 'STT', 6, Alignment(horizontal='center')),
        ('B7', 'Mã lỗi', 10, Alignment(horizontal='center')),
        ('C7', 'Tên nhóm lỗi', 38, Alignment(horizontal='left', indent=1)),
        ('D7', 'Số báo cáo', 14, Alignment(horizontal='right')),
        ('E7', 'Tỷ lệ (%)', 14, Alignment(horizontal='right')),
        ('F7', 'Lũy kế (%)', 14, Alignment(horizontal='right')),
        ('G7', 'Phân loại Pareto', 20, Alignment(horizontal='center'))
    ]

    for cell_id, title, width, align in headers:
        ws1[cell_id] = title
        ws1[cell_id].font = font_tbl_head
        ws1[cell_id].fill = fill_blue_head
        ws1[cell_id].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws1[cell_id].border = border_thin
        col_letter = cell_id[0]
        ws1.column_dimensions[col_letter].width = width

    data_rows = [
        (1, 'L02', 'Đặt hàng không thành công', 28, 'Vital Few (Ưu tiên tuần 5)'),
        (2, 'L06', 'Giỏ hàng tính sai tổng tiền', 22, 'Vital Few (Ưu tiên tuần 5)'),
        (3, 'L04', 'Tìm kiếm và lọc không chính xác', 16, 'Vital Few (Ưu tiên tuần 5)'),
        (4, 'L01', 'Cấu hình máy tính sai hoặc thiếu', 14, 'Vital Few (Mốc 80.0%)'),
        (5, 'L07', 'Hình ảnh sản phẩm bị lỗi', 8, 'Trivial Many (Thứ yếu)'),
        (6, 'L03', 'Không gửi email xác nhận', 5, 'Trivial Many (Thứ yếu)'),
        (7, 'L08', 'Trang sản phẩm tải chậm', 4, 'Trivial Many (Thứ yếu)'),
        (8, 'L05', 'Xem được đơn hàng của người khác', 3, 'Trivial Many (Hotfix bảo mật!)')
    ]

    start_row = 8
    for idx, (stt, code, name, count, category) in enumerate(data_rows):
        r = start_row + idx
        ws1.row_dimensions[r].height = 22
        
        ws1[f'A{r}'] = stt
        ws1[f'A{r}'].alignment = Alignment(horizontal='center', vertical='center')
        
        ws1[f'B{r}'] = code
        ws1[f'B{r}'].alignment = Alignment(horizontal='center', vertical='center')
        ws1[f'B{r}'].font = font_tbl_bold
        
        ws1[f'C{r}'] = name
        ws1[f'C{r}'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
        
        ws1[f'D{r}'] = count
        ws1[f'D{r}'].alignment = Alignment(horizontal='right', vertical='center')
        ws1[f'D{r}'].font = font_tbl_bold
        ws1[f'D{r}'].number_format = '#,##0'

        ws1[f'E{r}'] = f"=D{r}/$D$16"
        ws1[f'E{r}'].alignment = Alignment(horizontal='right', vertical='center')
        ws1[f'E{r}'].number_format = '0.0%'

        ws1[f'F{r}'] = f"=SUM($D$8:D{r})/$D$16"
        ws1[f'F{r}'].alignment = Alignment(horizontal='right', vertical='center')
        ws1[f'F{r}'].font = font_tbl_bold
        ws1[f'F{r}'].number_format = '0.0%'

        ws1[f'G{r}'] = category
        ws1[f'G{r}'].alignment = Alignment(horizontal='center', vertical='center')

        row_fill = fill_vital if idx < 4 else fill_trivial
        for col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
            cell = ws1[f'{col_letter}{r}']
            cell.font = font_tbl_bold if col_letter in ['B', 'D', 'F'] else font_tbl_cell
            cell.fill = row_fill
            cell.border = border_thin

    # Summary Row
    sum_row = start_row + len(data_rows)
    ws1.row_dimensions[sum_row].height = 24
    ws1[f'A{sum_row}'] = "Tổng cộng"
    ws1.merge_cells(f'A{sum_row}:C{sum_row}')
    ws1[f'A{sum_row}'].font = Font(name='Segoe UI', size=11, bold=True, color='1E3A8A')
    ws1[f'A{sum_row}'].alignment = Alignment(horizontal='center', vertical='center')

    ws1[f'D{sum_row}'] = f"=SUM(D{start_row}:D{sum_row-1})"
    ws1[f'D{sum_row}'].font = Font(name='Segoe UI', size=11, bold=True, color='1E3A8A')
    ws1[f'D{sum_row}'].alignment = Alignment(horizontal='right', vertical='center')
    ws1[f'D{sum_row}'].number_format = '#,##0'

    ws1[f'E{sum_row}'] = f"=SUM(E{start_row}:E{sum_row-1})"
    ws1[f'E{sum_row}'].font = Font(name='Segoe UI', size=11, bold=True, color='1E3A8A')
    ws1[f'E{sum_row}'].alignment = Alignment(horizontal='right', vertical='center')
    ws1[f'E{sum_row}'].number_format = '0.0%'

    ws1[f'F{sum_row}'] = f"=F{sum_row-1}"
    ws1[f'F{sum_row}'].font = Font(name='Segoe UI', size=11, bold=True, color='1E3A8A')
    ws1[f'F{sum_row}'].alignment = Alignment(horizontal='right', vertical='center')
    ws1[f'F{sum_row}'].number_format = '0.0%'

    ws1[f'G{sum_row}'] = "100.0% kiểm tra khớp"
    ws1[f'G{sum_row}'].font = Font(name='Segoe UI', size=10, italic=True, color='475569')
    ws1[f'G{sum_row}'].alignment = Alignment(horizontal='center', vertical='center')

    for col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
        cell = ws1[f'{col_letter}{sum_row}']
        cell.fill = fill_total
        cell.border = border_total

    # Embed Chart Image into Excel
    img_chart1 = openpyxl.drawing.image.Image('public/images/pareto_baitap_canhan.png')
    img_chart1.width = 750
    img_chart1.height = 468
    ws1.add_image(img_chart1, 'A18')

    # Questions & Answers Section (Below Chart)
    q_start = 43
    ws1[f'A{q_start}'] = "2. KẾT LUẬN VÀ TRẢ LỜI CÂU HỎI BÀI THỰC HÀNH CÁ NHÂN"
    ws1.merge_cells(f'A{q_start}:G{q_start}')
    ws1[f'A{q_start}'].font = font_sec_head
    ws1[f'A{q_start}'].fill = fill_navy
    ws1[f'A{q_start}'].alignment = Alignment(vertical='center', indent=1)

    qa_list = [
        ("Câu 1: Cần tối thiểu bao nhiêu nhóm lỗi đầu tiên để đạt ít nhất 80%? Ghi mã lỗi và tỷ lệ lũy kế.",
         "Trả lời:\n"
         "• Cần tối thiểu 4 nhóm lỗi đầu tiên để đạt đúng ngưỡng 80% tổng số báo cáo lỗi.\n"
         "• Danh sách 4 mã lỗi và tỷ lệ lũy kế tương ứng:\n"
         "   1. L02 (Đặt hàng không thành công): 28 báo cáo | Tỷ lệ: 28.0% | Lũy kế: 28.0%\n"
         "   2. L06 (Giỏ hàng tính sai tổng tiền): 22 báo cáo | Tỷ lệ: 22.0% | Lũy kế: 50.0%\n"
         "   3. L04 (Tìm kiếm và lọc không chính xác): 16 báo cáo | Tỷ lệ: 16.0% | Lũy kế: 66.0%\n"
         "   4. L01 (Cấu hình máy tính sai hoặc thiếu): 14 báo cáo | Tỷ lệ: 14.0% | Lũy kế: 80.0% (Đạt mốc 80.0%)."),

        ("Câu 2: Các nhóm được chọn chiếm bao nhiêu phần trăm trong 8 nhóm? Dữ liệu có đúng tỷ lệ 80–20 không?",
         "Trả lời:\n"
         "• Tỷ lệ các nhóm được chọn trong tổng số nhóm: 4 / 8 nhóm = 50.0%.\n"
         "• Nhận xét về quy tắc 80–20:\n"
         "   + Tỷ lệ thực tế phân tích từ dữ liệu dự án là 80–50 (50% số nhóm nguyên nhân gây ra 80% tổng số báo cáo lỗi), chưa khớp tuyệt đối với con số lý thuyết kinh điển 80–20 (theo lý thuyết chuẩn 20% tương ứng khoảng 1.6 ~ 2 nhóm lỗi).\n"
         "   + Tuy nhiên, dữ liệu vẫn tuân thủ chặt chẽ Nguyên lý Pareto trong Quản lý dự án thực tế: Phân định rõ nhóm 'Vital Few' (4 lỗi đầu dòng tập trung phần lớn vấn đề cốt lõi liên quan trực tiếp đến luồng đặt hàng và tính tiền của website) giúp PM/Team ưu tiên tập trung nguồn lực trong Tuần 5 giải quyết dứt điểm 4 nhóm này để cải thiện 80% chất lượng hệ thống thay vì phân tán dàn trải."),

        ("Câu 3: L05 là lỗi phân quyền nghiêm trọng. Có nên trì hoãn xử lý vì ít báo cáo không? Giải thích ngắn gọn.",
         "Trả lời: KHÔNG ĐƯỢC PHÉP TRÌ HOÃN, MÀ PHẢI XỬ LÝ KHẨN CẤP NGAY LẬP TỨC (Mức độ ưu tiên Critical / Hotfix P1).\n"
         "• Giải thích ngắn gọn:\n"
         "   1. Hạn chế của Biểu đồ Pareto: Sơ đồ Pareto phân loại theo TẦN SUẤT XUẤT HIỆN (Frequency/Count) chứ không phản ánh MỨC ĐỘ NGUY HIỂM / TÁC ĐỘNG NGHIỆM TRỌNG (Severity/Impact).\n"
         "   2. Tính chất nguy hiểm của lỗi L05: 'Xem được đơn hàng của người khác' là lỗ hổng an ninh phân quyền truy cập nghiêm trọng (IDOR - Insecure Direct Object References theo chuẩn OWASP). Lỗ hổng này làm rò rỉ dữ liệu cá nhân khách hàng (PII - tên, địa chỉ, số điện thoại, thông tin mua hàng), vi phạm pháp luật bảo vệ dữ liệu cá nhân, có nguy cơ bị tấn công dữ liệu hàng loạt và gây khủng hoảng truyền thông phá hủy uy tín doanh nghiệp.\n"
         "   3. Chiến lược Quản lý chất lượng & Rủi ro PMBOK: Đội ngũ dự án phải áp dụng đồng thời Ma trận Mức độ nghiêm trọng lỗi. L05 dù chỉ có 3 báo cáo nhưng thuộc nhóm rủi ro cao nhất, bắt buộc phải phát hành bản vá nóng (Emergency Hotfix) ngay trong ngày, song song với việc khắc phục nhóm Vital Few (L02, L06, L04, L01) trong Tuần 5.")
    ]

    curr_row = q_start + 1
    for q_title, q_body in qa_list:
        ws1[f'A{curr_row}'] = q_title
        ws1.merge_cells(f'A{curr_row}:G{curr_row}')
        ws1[f'A{curr_row}'].font = font_q_head
        ws1[f'A{curr_row}'].fill = fill_q_box
        ws1.row_dimensions[curr_row].height = 24

        curr_row += 1
        ws1[f'A{curr_row}'] = q_body
        ws1.merge_cells(f'A{curr_row}:G{curr_row}')
        ws1[f'A{curr_row}'].font = font_ans
        ws1[f'A{curr_row}'].alignment = Alignment(wrap_text=True, vertical='top')
        lines = q_body.count('\n') + 1
        ws1.row_dimensions[curr_row].height = max(55, lines * 19)

        curr_row += 1


    # --------------------------------------------------------------------------
    # SHEET 2: DỰ ÁN THỰC TẾ GODDY RECRUIT (Phần mềm quản lý hóa đơn & công nợ B2B)
    # --------------------------------------------------------------------------
    ws2 = wb.create_sheet(title="Pareto Dự Án GODDY Recruit")
    ws2.views.sheetView[0].showGridLines = True

    ws2['A1'] = "PHÂN TÍCH PARETO CHẤT LƯỢNG HỆ THỐNG - DỰ ÁN GODDY RECRUIT"
    ws2['A1'].font = font_title
    ws2['A2'] = "Đề tài: Quản Lý Hóa Đơn & Công Nợ Tích Hợp Dashboard Dữ Liệu Tuyển Dụng B2B (Nhóm 10)"
    ws2['A2'].font = font_subtitle

    ws2['A4'] = "Dự án:"
    ws2['A4'].font = font_meta
    ws2['B4'] = "GODDY RECRUIT (Nhóm 10 - Quản lý dự án CNTT)"
    ws2['B4'].font = font_meta_val

    ws2['D4'] = "Giai đoạn kiểm thử:"
    ws2['D4'].font = font_meta
    ws2['E4'] = "Sprint 1 - 4 (System & UAT Testing)"
    ws2['E4'].font = font_meta_val

    ws2['A6'] = "1. BẢNG TỔNG HỢP LỖI KIỂM THỬ VÀ PHÂN TÍCH VITAL FEW THEO PARETO"
    ws2.merge_cells('A6:G6')
    ws2['A6'].font = font_sec_head
    ws2['A6'].fill = fill_navy
    ws2['A6'].alignment = Alignment(vertical='center', indent=1)

    for cell_id, title, width, align in headers:
        ws2[cell_id] = title
        ws2[cell_id].font = font_tbl_head
        ws2[cell_id].fill = fill_blue_head
        ws2[cell_id].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws2[cell_id].border = border_thin
        col_letter = cell_id[0]
        ws2.column_dimensions[col_letter].width = width

    project_data = [
        (1, 'L01', 'Lệch tính tiền thuế VAT 8% & làm tròn số tiền', 48, 'Vital Few (40.0% tổng lỗi)'),
        (2, 'L02', 'Lệch đồng bộ trạng thái công nợ khi trả dần (Paid/Partial)', 36, 'Vital Few (Lũy kế 70.0%)'),
        (3, 'L03', 'Tính sai ngày hết hạn bảo hành ứng viên 60 ngày', 14, 'Vital Few (Lũy kế 81.7% - Đạt mốc)'),
        (4, 'L04', 'Trùng mã số thuế hoặc thiếu contact B2B', 8, 'Trivial Many (Lũy kế 88.3%)'),
        (5, 'L05', 'Lỗi định dạng font UTF-8 tiếng Việt in hóa đơn điện tử', 6, 'Trivial Many (Lũy kế 93.3%)'),
        (6, 'L06', 'Báo cáo Aging Report tải chậm khi nhiều hóa đơn', 4, 'Trivial Many (Lũy kế 96.7%)'),
        (7, 'L07', 'Lỗi phân quyền truy cập chéo hóa đơn đối tác B2B', 3, 'Trivial Many (Lỗ hổng bảo mật P1!)'),
        (8, 'L08', 'Lệch cấu trúc bảng khi xuất báo cáo CSV', 1, 'Trivial Many (Lũy kế 100.0%)')
    ]

    for idx, (stt, code, name, count, category) in enumerate(project_data):
        r = start_row + idx
        ws2.row_dimensions[r].height = 22
        
        ws2[f'A{r}'] = stt
        ws2[f'A{r}'].alignment = Alignment(horizontal='center', vertical='center')
        
        ws2[f'B{r}'] = code
        ws2[f'B{r}'].alignment = Alignment(horizontal='center', vertical='center')
        ws2[f'B{r}'].font = font_tbl_bold
        
        ws2[f'C{r}'] = name
        ws2[f'C{r}'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
        
        ws2[f'D{r}'] = count
        ws2[f'D{r}'].alignment = Alignment(horizontal='right', vertical='center')
        ws2[f'D{r}'].font = font_tbl_bold
        ws2[f'D{r}'].number_format = '#,##0'

        ws2[f'E{r}'] = f"=D{r}/$D$16"
        ws2[f'E{r}'].alignment = Alignment(horizontal='right', vertical='center')
        ws2[f'E{r}'].number_format = '0.0%'

        ws2[f'F{r}'] = f"=SUM($D$8:D{r})/$D$16"
        ws2[f'F{r}'].alignment = Alignment(horizontal='right', vertical='center')
        ws2[f'F{r}'].font = font_tbl_bold
        ws2[f'F{r}'].number_format = '0.0%'

        ws2[f'G{r}'] = category
        ws2[f'G{r}'].alignment = Alignment(horizontal='center', vertical='center')

        row_fill = fill_vital if idx < 3 else fill_trivial
        for col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
            cell = ws2[f'{col_letter}{r}']
            cell.font = font_tbl_bold if col_letter in ['B', 'D', 'F'] else font_tbl_cell
            cell.fill = row_fill
            cell.border = border_thin

    # Summary Row
    ws2.row_dimensions[sum_row].height = 24
    ws2[f'A{sum_row}'] = "Tổng cộng"
    ws2.merge_cells(f'A{sum_row}:C{sum_row}')
    ws2[f'A{sum_row}'].font = Font(name='Segoe UI', size=11, bold=True, color='1E3A8A')
    ws2[f'A{sum_row}'].alignment = Alignment(horizontal='center', vertical='center')

    ws2[f'D{sum_row}'] = f"=SUM(D{start_row}:D{sum_row-1})"
    ws2[f'D{sum_row}'].font = Font(name='Segoe UI', size=11, bold=True, color='1E3A8A')
    ws2[f'D{sum_row}'].alignment = Alignment(horizontal='right', vertical='center')
    ws2[f'D{sum_row}'].number_format = '#,##0'

    ws2[f'E{sum_row}'] = f"=SUM(E{start_row}:E{sum_row-1})"
    ws2[f'E{sum_row}'].font = Font(name='Segoe UI', size=11, bold=True, color='1E3A8A')
    ws2[f'E{sum_row}'].alignment = Alignment(horizontal='right', vertical='center')
    ws2[f'E{sum_row}'].number_format = '0.0%'

    ws2[f'F{sum_row}'] = f"=F{sum_row-1}"
    ws2[f'F{sum_row}'].font = Font(name='Segoe UI', size=11, bold=True, color='1E3A8A')
    ws2[f'F{sum_row}'].alignment = Alignment(horizontal='right', vertical='center')
    ws2[f'F{sum_row}'].number_format = '0.0%'

    ws2[f'G{sum_row}'] = "100.0% kiểm tra khớp"
    ws2[f'G{sum_row}'].font = Font(name='Segoe UI', size=10, italic=True, color='475569')
    ws2[f'G{sum_row}'].alignment = Alignment(horizontal='center', vertical='center')

    for col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
        cell = ws2[f'{col_letter}{sum_row}']
        cell.fill = fill_total
        cell.border = border_total

    # Embed Chart Image
    img_chart2 = openpyxl.drawing.image.Image('public/images/pareto_du_an_goddy.png')
    img_chart2.width = 750
    img_chart2.height = 468
    ws2.add_image(img_chart2, 'A18')

    # Project Evaluation & Action Plan
    q_start = 43
    ws2[f'A{q_start}'] = "2. PHÂN TÍCH CHUYÊN SÂU & KẾ HOẠCH BẢO ĐẢM CHẤT LƯỢNG (QA/QC) CHO DỰ ÁN GODDY"
    ws2.merge_cells(f'A{q_start}:G{q_start}')
    ws2[f'A{q_start}'].font = font_sec_head
    ws2[f'A{q_start}'].fill = fill_navy
    ws2[f'A{q_start}'].alignment = Alignment(vertical='center', indent=1)

    project_qa = [
        ("Phân định Nhóm 'Vital Few' của Dự án GODDY Recruit (Quy tắc 80/20):",
         "• Có 3 nhóm lỗi đầu tiên (L01, L02, L03) chiếm 3/8 = 37.5% tổng số nhóm nhưng gây ra tới 81.7% tổng số lỗi phát sinh trong hệ thống.\n"
         "• Bản chất nghiệp vụ của Vital Few trong dự án GODDY:\n"
         "   + L01 (Lệch thuế VAT 8% - 40%): Là lỗi tài chính kế toán trực tiếp, gây sai hóa đơn xuất cho doanh nghiệp đối tác (FPT, VNG, Shopee).\n"
         "   + L02 (Lệch trạng thái công nợ - 30%): Gây sai lệch báo cáo tài chính và tỷ lệ thu hồi nợ trên Dashboard.\n"
         "   + L03 (Sai hạn bảo hành 60 ngày - 11.7%): Dẫn đến tranh chấp bồi thường deal tuyển dụng khi ứng viên nghỉ việc.\n"
         "==> KẾ HOẠCH HÀNH ĐỘNG: Tập trung 80% nguồn lực kiểm thử tự động (tests/api.test.js) và refactor code backend cho 3 module này trước tiên."),

        ("Đánh giá rủi ro an ninh & ngoại lệ Pareto đối với lỗi L07 (Phân quyền B2B):",
         "• L07 chỉ ghi nhận 3 báo cáo (2.5% - nằm trong Trivial Many) nhưng là lỗ hổng bảo mật nghiêm trọng (IDOR - Insecure Direct Object References).\n"
         "• Nếu doanh nghiệp A (ví dụ FPT) có thể xem trộm hóa đơn hoặc mức lương ứng viên của doanh nghiệp B (Shopee), dự án sẽ vi phạm nghiêm trọng hợp đồng bảo mật (NDA).\n"
         "• BIỆN PHÁP XỬ LÝ: Nhóm 10 kích hoạt quy trình Emergency Hotfix ngay lập tức, bổ sung middleware phân quyền chặt chẽ `requireRole('client')` kèm điều kiện `WHERE clientId = req.user.clientId` trong Sequelize để vá triệt để."),

        ("Kế hoạch kiểm soát chất lượng (Quality Control) trong các Sprint tiếp theo:",
         "• Tích hợp bộ Unit Test tự động (Jest/Mocha/Node Test) chạy tự động qua lệnh `npm test` trước mỗi lần Git Commit.\n"
         "• Thiết lập kiểm soát giao dịch (Database Transaction) trong Sequelize để đảm bảo tính toàn vẹn (ACID) khi ghi nhận thanh toán trừ dần nợ.\n"
         "• Đưa biểu đồ Pareto trực quan vào hệ thống Dashboard Quản trị để theo dõi chất lượng phần mềm liên tục theo thời gian thực.")
    ]

    curr_row = q_start + 1
    for q_title, q_body in project_qa:
        ws2[f'A{curr_row}'] = q_title
        ws2.merge_cells(f'A{curr_row}:G{curr_row}')
        ws2[f'A{curr_row}'].font = font_q_head
        ws2[f'A{curr_row}'].fill = fill_q_box
        ws2.row_dimensions[curr_row].height = 24

        curr_row += 1
        ws2[f'A{curr_row}'] = q_body
        ws2.merge_cells(f'A{curr_row}:G{curr_row}')
        ws2[f'A{curr_row}'].font = font_ans
        ws2[f'A{curr_row}'].alignment = Alignment(wrap_text=True, vertical='top')
        lines = q_body.count('\n') + 1
        ws2.row_dimensions[curr_row].height = max(55, lines * 19)

        curr_row += 1

    try:
        wb.save('MSSV_HoTen_Pareto.xlsx')
        print('-> Đã tạo file Excel nộp bài: MSSV_HoTen_Pareto.xlsx')
    except Exception as e:
        print(f'Lưu ý MSSV_HoTen_Pareto.xlsx: {e}')

    try:
        wb.save('HuynhNguyenVinhPhuc_Pareto.xlsx')
        print('-> Đã tạo file Excel nộp bài: HuynhNguyenVinhPhuc_Pareto.xlsx')
    except Exception as e:
        wb.save('HuynhNguyenVinhPhuc_Pareto_v2.xlsx')
        print(f'-> File HuynhNguyenVinhPhuc_Pareto.xlsx đang mở trong Excel, đã lưu thành: HuynhNguyenVinhPhuc_Pareto_v2.xlsx')

if __name__ == '__main__':
    generate_practice_chart()
    generate_project_chart()
    generate_excel_file()
    print('=== TẤT CẢ CÁC ARTIFACTS ĐÃ ĐƯỢC TẠO THÀNH CÔNG ===')
